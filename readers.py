"""Turn files and web pages into text 'pages' with source labels (PDF, Word, PowerPoint, Excel, CSV, HTML, text, images, URLs)."""
from __future__ import annotations

import io
import ipaddress
import re
import socket
from dataclasses import dataclass, field
from urllib.parse import urldefrag, urljoin, urlparse

import httpx
import pandas as pd
from PIL import Image, ImageOps

MAX_CHARS = 1_500_000      # per document
MAX_BYTES = 20 * 1024 * 1024
OCR_PAGE_LIMIT = 25        # scanned PDF pages read with OCR per file
HEADERS = {"User-Agent": "DocChat/1.0 (student project)"}
TEXT_EXT = {"txt", "md", "markdown", "json", "log", "xml", "yaml", "yml", "rst", "py", "js", "java", "c", "cpp", "sql", "ini", "tex"}
IMAGE_EXT = {"png", "jpg", "jpeg", "webp", "bmp", "tif", "tiff"}
SUPPORTED = sorted({"pdf", "docx", "pptx", "xlsx", "xls", "csv", "html", "htm"} | TEXT_EXT | IMAGE_EXT)


@dataclass
class Document:
    name: str
    kind: str
    pages: list                      # [(label, text)]
    notes: list = field(default_factory=list)

    @property
    def text(self) -> str:
        return "\n\n".join(t for _, t in self.pages)

    @property
    def chars(self) -> int:
        return sum(len(t) for _, t in self.pages)


# ───────────── OCR for images and scanned pages ─────────────
def _rapid_text(img: Image.Image) -> str:
    try:
        import numpy as np
        from rapidocr_onnxruntime import RapidOCR
        res, _ = RapidOCR()(np.array(img.convert("RGB"))[:, :, ::-1])
        return "\n".join(str(r[1]) for r in (res or []))
    except Exception:
        return ""


def ocr_image(img: Image.Image, lang: str = "eng") -> str:
    """Tesseract first; RapidOCR (handwriting-friendly) if Tesseract finds almost nothing."""
    text = ""
    try:
        import pytesseract
        have = set(pytesseract.get_languages(config=""))
        lang = "+".join(l for l in lang.split("+") if l in have) or "eng"
        g = ImageOps.autocontrast(ImageOps.exif_transpose(img).convert("L"))
        longest = max(g.size)
        if longest < 1600:
            f = min(1600 / longest, 3.0)
            g = g.resize((int(g.width * f), int(g.height * f)), Image.LANCZOS)
        text = pytesseract.image_to_string(g, lang=lang, config="--oem 1 --psm 3")
    except Exception:
        pass
    if len(text.strip()) < 20:
        text = _rapid_text(img) or text
    return text.strip()


# ───────────── per-format readers ─────────────
def _pdf(name: str, raw: bytes, lang: str) -> Document:
    try:
        import pymupdf as fitz
    except ImportError:
        import fitz
    pages, scanned, skipped = [], 0, 0
    with fitz.open(stream=raw, filetype="pdf") as doc:
        if doc.needs_pass:
            raise ValueError("This PDF is password protected.")
        for i, page in enumerate(doc, 1):
            text = page.get_text("text").strip()
            if len(text) < 30:  # probably a scan
                if scanned < OCR_PAGE_LIMIT:
                    pix = page.get_pixmap(dpi=150, alpha=False)
                    text = ocr_image(Image.frombytes("RGB", (pix.width, pix.height), pix.samples), lang)
                    scanned += 1
                else:
                    skipped += 1
            if text:
                pages.append((f"p. {i}", text))
    notes = []
    if scanned:
        notes.append(f"{scanned} scanned page(s) were read with OCR")
    if skipped:
        notes.append(f"{skipped} scanned page(s) skipped (OCR limit is {OCR_PAGE_LIMIT} per file)")
    return Document(name, "pdf", pages, notes)


def _docx(name: str, raw: bytes) -> Document:
    from docx import Document as Docx
    d = Docx(io.BytesIO(raw))
    parts = []
    for p in d.paragraphs:
        t = p.text.strip()
        if t:
            style = p.style.name if p.style is not None else ""
            parts.append(f"# {t}" if style.startswith("Heading") or style == "Title" else t)
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return Document(name, "docx", [("", "\n".join(parts))])


def _pptx(name: str, raw: bytes) -> Document:
    from pptx import Presentation
    pages = []
    for i, slide in enumerate(Presentation(io.BytesIO(raw)).slides, 1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            parts.append("Notes: " + slide.notes_slide.notes_text_frame.text.strip())
        if parts:
            pages.append((f"slide {i}", "\n".join(parts)))
    return Document(name, "pptx", pages)


def _table_text(df: pd.DataFrame) -> str:
    df = df.dropna(how="all").head(3000)
    lines = []
    for _, row in df.iterrows():
        cells = [f"{c}: {v}" for c, v in row.items() if pd.notna(v) and str(v).strip()]
        if cells:
            lines.append("; ".join(cells))
    return "\n".join(lines)


def _excel(name: str, raw: bytes) -> Document:
    sheets = pd.read_excel(io.BytesIO(raw), sheet_name=None)
    return Document(name, "xlsx", [(f"sheet {s}", _table_text(df)) for s, df in sheets.items() if len(df)])


def _csv(name: str, raw: bytes) -> Document:
    df = pd.read_csv(io.BytesIO(raw), sep=None, engine="python", nrows=5000, encoding_errors="replace")
    return Document(name, "csv", [("", _table_text(df))], [f"read the first {min(len(df), 5000)} rows"] if len(df) >= 5000 else [])


def html_to_text(html: str) -> tuple[str, str]:
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(html, "html.parser")
    title = soup.title.get_text(" ", strip=True) if soup.title else ""
    for t in soup(["script", "style", "noscript", "nav", "footer", "form", "svg", "iframe", "template"]):
        t.decompose()
    text = (soup.body or soup).get_text("\n", strip=True)
    return title, re.sub(r"\n{3,}", "\n\n", text)


def _html(name: str, html: str) -> Document:
    title, text = html_to_text(html)
    return Document(name, "html", [(title, text)] if text else [])


def _text(name: str, raw: bytes, kind: str = "text") -> Document:
    return Document(name, kind, [("", raw.decode("utf-8", errors="replace"))])


def _image(name: str, raw: bytes, lang: str) -> Document:
    img = Image.open(io.BytesIO(raw))
    img.load()
    text = ocr_image(img, lang)
    return Document(name, "image", [("", text)] if text else [], [] if text else ["no readable text was found in this image"])


def read_file(name: str, raw: bytes, lang: str = "eng") -> Document:
    """Read any supported file into a Document."""
    ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
    if raw[:5] == b"%PDF-" or ext == "pdf":
        doc = _pdf(name, raw, lang)
    elif ext == "docx":
        doc = _docx(name, raw)
    elif ext == "pptx":
        doc = _pptx(name, raw)
    elif ext in ("xlsx", "xls"):
        doc = _excel(name, raw)
    elif ext == "csv":
        doc = _csv(name, raw)
    elif ext in ("html", "htm"):
        doc = _html(name, raw.decode("utf-8", errors="replace"))
    elif ext in IMAGE_EXT:
        doc = _image(name, raw, lang)
    elif ext in TEXT_EXT or ext == "":
        doc = _text(name, raw)
    else:
        raise ValueError(f"'.{ext}' files are not supported. Supported: {', '.join(SUPPORTED)}")
    doc.pages = [(l, t) for l, t in doc.pages if t.strip()]
    total = 0
    for i, (label, t) in enumerate(doc.pages):  # cap very large documents
        if total + len(t) > MAX_CHARS:
            doc.pages = doc.pages[:i] + [(label, t[: max(0, MAX_CHARS - total)])]
            doc.notes.append("document was very large, so only the first part was read")
            break
        total += len(t)
    if not doc.pages:
        raise ValueError(f"No readable text found in '{name}'.")
    return doc


# ───────────── websites (with protection against private-network access) ─────────────
def check_url(url: str) -> None:
    u = urlparse(url)
    if u.scheme not in ("http", "https") or not u.hostname:
        raise ValueError("Enter a full web address starting with http:// or https://")
    try:
        infos = socket.getaddrinfo(u.hostname, u.port or (443 if u.scheme == "https" else 80), proto=socket.IPPROTO_TCP)
    except socket.gaierror:
        raise ValueError(f"Could not find the website '{u.hostname}'.")
    for info in infos:
        ip = ipaddress.ip_address(info[4][0])
        if ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_reserved or ip.is_multicast or ip.is_unspecified:
            raise ValueError("That address points to a private network and is blocked for safety.")


def _fetch(url: str) -> tuple[str, str, bytes]:
    """GET with manual redirects (each hop re-checked). Returns (final_url, content_type, body)."""
    with httpx.Client(headers=HEADERS, timeout=20, follow_redirects=False) as c:
        for _ in range(5):
            check_url(url)
            with c.stream("GET", url) as r:
                if r.is_redirect:
                    url = urljoin(url, r.headers.get("location", ""))
                    continue
                r.raise_for_status()
                body, size = [], 0
                for part in r.iter_bytes():
                    body.append(part)
                    size += len(part)
                    if size > MAX_BYTES:
                        break
                return url, r.headers.get("content-type", "").lower(), b"".join(body)
    raise ValueError("Too many redirects.")


def _links(html: str, base: str) -> list[str]:
    from bs4 import BeautifulSoup
    host, out = urlparse(base).netloc, []
    for a in BeautifulSoup(html, "html.parser").find_all("a", href=True):
        u = urldefrag(urljoin(base, a["href"]))[0]
        p = urlparse(u)
        if p.scheme in ("http", "https") and p.netloc == host and u not in out \
                and not re.search(r"\.(png|jpe?g|gif|svg|css|js|zip|mp4|mp3|webp)$", p.path, re.I):
            out.append(u)
    return out


def read_url(url: str, lang: str = "eng", max_pages: int = 1) -> Document:
    """Read a web page (or PDF link). With max_pages > 1, follow links on the same site."""
    url = url.strip()
    if not re.match(r"^https?://", url):
        url = "https://" + url
    host = urlparse(url).netloc
    queue, seen, pages, notes = [url], set(), [], []
    while queue and len(seen) < max_pages:
        u = queue.pop(0)
        if u in seen:
            continue
        seen.add(u)
        try:
            final, ctype, raw = _fetch(u)
        except Exception as e:
            if not seen - {u}:
                raise ValueError(f"Could not open the page: {e}")
            notes.append(f"skipped {u}")
            continue
        if "pdf" in ctype or raw[:5] == b"%PDF-":
            d = _pdf(final, raw, lang)
            pages += [(f"{urlparse(final).path or '/'} {lbl}", t) for lbl, t in d.pages]
        elif "html" in ctype or b"<html" in raw[:2000].lower():
            html = raw.decode("utf-8", errors="replace")
            title, text = html_to_text(html)
            if text:
                pages.append((title or urlparse(final).path or "/", text))
            if max_pages > 1:
                queue += [l for l in _links(html, final) if l not in seen]
        else:
            pages.append((urlparse(final).path or "/", raw.decode("utf-8", errors="replace")))
    if not pages:
        raise ValueError("No readable text found at that address.")
    notes.append(f"read {len(seen)} page(s)")
    return Document(host + (urlparse(url).path.rstrip("/") if urlparse(url).path != "/" else ""), "website", pages, notes)
