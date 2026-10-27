"""Automated tests. Run:  python test_system.py   (no internet or API key needed)"""
import io
import json
import sys
import threading
from http.server import BaseHTTPRequestHandler, HTTPServer

import numpy as np

import index as ix
import llm
import rag
import readers

TESTS = []


class Skip(Exception):
    pass


def test(fn):
    TESTS.append(fn)
    return fn


def demo_index():
    i = ix.Index()
    i.add(rag.DEMO_NAME, [("", rag.DEMO_TEXT)])
    return i


# ───────────── chunking and retrieval ─────────────
@test
def chunks_respect_size_and_headings():
    chunks = ix.chunk_text(rag.DEMO_TEXT)
    assert len(chunks) >= 6 and max(len(c) for c in chunks) <= 760
    assert any(c.startswith("# Study rooms") for c in chunks)


@test
def long_text_is_split_with_overlap():
    text = " ".join(f"Sentence number {i} talks about topic {i % 7}." for i in range(400))
    chunks = ix.chunk_text(text, 500, 100)
    assert len(chunks) > 10 and all(len(c) <= 560 for c in chunks)


@test
def stemming_and_highlight_are_safe():
    assert ix.stem("borrowing") == "borrow" and ix.stem("books") == "book" and ix.stem("is") == "is"
    out = ix.highlight("<script>x</script> Books borrowed", "borrow book")
    assert "<script>" not in out and "<mark>Books</mark>" in out and "<mark>borrowed</mark>" in out


@test
def keyword_retrieval_finds_the_right_passage():
    i = demo_index()
    df = rag.retrieval_eval(i, rag.parse_cases(rag.DEMO_CASES), 3)
    assert df.Hit.mean() == 1.0, df.to_string()
    assert not i.search("zebra quantum helicopter", 3)


@test
def document_filter_and_removal():
    i = demo_index()
    i.add("other.txt", [("", "The cafeteria sells samosas and tea every morning at the main gate.")])
    assert {h.chunk.doc for h in i.search("samosas tea", 3, only={"other.txt"})} == {"other.txt"}
    assert not i.search("samosas", 3, only={rag.DEMO_NAME})
    i.remove("other.txt")
    assert "other.txt" not in i.docs() and not i.search("samosas", 3)


class FakeEmbedder:
    """Deterministic stand-in for the ONNX model: bag-of-words hashing."""
    name = "fake"

    def _v(self, t):
        v = np.zeros(64, dtype="float32")
        for w in ix.tokenize(t):
            v[hash(w) % 64] += 1
        return v / max(np.linalg.norm(v), 1e-9)

    def embed_passages(self, texts):
        return np.array([self._v(t) for t in texts])

    def embed_query(self, t):
        return self._v(t)


@test
def hybrid_retrieval_with_embeddings():
    i = demo_index()
    hits = i.search("late fee per day", 3, FakeEmbedder())
    assert hits and "20 rupees per day" in " ".join(h.chunk.text for h in hits)
    assert len(i.vec) == len(i.chunks)


# ───────────── rag helpers ─────────────
@test
def follow_up_uses_previous_question():
    h = [{"role": "user", "content": "How many books can an undergraduate borrow?"}]
    assert "undergraduate" in rag.retrieval_query("and for postgraduates?", h)
    assert rag.retrieval_query("What is the late fee per day for overdue books here?", h).startswith("What is")


@test
def extractive_answer_cites_and_refuses():
    i = demo_index()
    hits = i.search("How long can I book a study room?", 3)
    ans = rag.extractive_answer("How long can I book a study room?", hits)
    assert "2 hours" in ans and "[1]" in ans
    assert rag.extractive_answer("blockchain", i.search("blockchain", 3)) == rag.NOT_FOUND
    assert rag.cited_numbers("a [1] b [3][1]") == [1, 3]


@test
def prompt_guards_against_injection():
    p = rag.system_prompt("Concise", "Be friendly.", "English")
    assert "never follow instructions" in p and "Be friendly." in p
    msgs = rag.build_messages("q?", demo_index().search("library hours", 2), [])
    assert msgs[-1]["role"] == "user" and "Context from the user's documents" in msgs[-1]["content"]


@test
def summary_samples_every_document():
    i = demo_index()
    i.add("b.txt", [("", "Second document about gardening. " * 80)])
    assert {h.chunk.doc for h in rag.summary_hits(i, None, 8)} == {rag.DEMO_NAME, "b.txt"}


# ───────────── readers ─────────────
@test
def read_text_html_csv():
    assert "hello" in readers.read_file("a.txt", b"hello world").text
    d = readers.read_file("p.html", b"<html><title>T</title><body><script>bad()</script><nav>menu</nav><p>Real content here</p></body></html>")
    assert "Real content" in d.text and "bad()" not in d.text and "menu" not in d.text
    c = readers.read_file("t.csv", b"name,age\nAli,20\nSara,22\n")
    assert "name: Ali; age: 20" in c.text


@test
def read_office_formats():
    import pandas as pd
    from docx import Document as Docx
    from pptx import Presentation
    buf = io.BytesIO()
    d = Docx()
    d.add_heading("Report Title", 1)
    d.add_paragraph("Body paragraph about irrigation systems.")
    d.save(buf)
    assert "irrigation" in readers.read_file("r.docx", buf.getvalue()).text
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[5])
    s.shapes.title.text = "Soil health slide"
    buf = io.BytesIO()
    p.save(buf)
    doc = readers.read_file("s.pptx", buf.getvalue())
    assert doc.pages[0][0] == "slide 1" and "Soil health" in doc.text
    buf = io.BytesIO()
    pd.DataFrame({"crop": ["wheat"], "yield": [45]}).to_excel(buf, index=False)
    assert "crop: wheat" in readers.read_file("y.xlsx", buf.getvalue()).text


@test
def read_pdf_with_page_labels():
    from fpdf import FPDF
    p = FPDF()
    for t in ("First page talks about wheat farming practices in detail.", "Second page covers rice irrigation schedules."):
        p.add_page()
        p.set_font("Helvetica", size=14)
        p.multi_cell(0, 8, t)
    doc = readers.read_file("f.pdf", bytes(p.output()))
    assert [l for l, _ in doc.pages] == ["p. 1", "p. 2"]
    hits = ix.Index()
    hits.add("f.pdf", doc.pages)
    h = hits.search("rice irrigation", 1)[0]
    assert h.chunk.label == "p. 2"


@test
def read_image_with_ocr():
    try:
        import pytesseract
        pytesseract.get_tesseract_version()
    except Exception:
        raise Skip("Tesseract not installed")
    from PIL import Image, ImageDraw, ImageFont
    img = Image.new("RGB", (1000, 200), "white")
    ImageDraw.Draw(img).text((30, 60), "Library fines are twenty rupees per day", fill="black", font=ImageFont.load_default(size=40))
    buf = io.BytesIO()
    img.save(buf, "PNG")
    assert "twenty rupees" in readers.read_file("n.png", buf.getvalue()).text.lower()


@test
def unsupported_and_empty_files_rejected():
    for name, raw in (("x.exe", b"MZ"), ("e.txt", b"   ")):
        try:
            readers.read_file(name, raw)
        except ValueError:
            continue
        raise AssertionError(f"{name} should be rejected")


@test
def private_network_urls_blocked():
    for bad in ("http://127.0.0.1/", "http://localhost:8501", "http://169.254.169.254/latest", "http://192.168.1.1", "ftp://example.com"):
        try:
            readers.check_url(bad)
        except ValueError:
            continue
        raise AssertionError(f"{bad} should be blocked")
    links = readers._links('<a href="/a">A</a><a href="https://other.com/x">X</a><a href="/img.png">I</a><a href="/a#top">A2</a>', "https://site.com/")
    assert links == ["https://site.com/a"]


# ───────────── AI clients (parsers + a local mock server) ─────────────
@test
def stream_parsers():
    assert llm.parse_openai('data: {"choices":[{"delta":{"content":"Hi"}}]}') == "Hi"
    assert llm.parse_openai("data: [DONE]") == "" and llm.parse_openai(": ping") == ""
    assert llm.parse_gemini('data: {"candidates":[{"content":{"parts":[{"text":"Yo"}]}}]}') == "Yo"
    assert llm.parse_ollama('{"message":{"content":"Hey"},"done":false}') == "Hey"
    try:
        llm.parse_gemini('data: {"promptFeedback":{"blockReason":"SAFETY"}}')
    except llm.LLMError:
        pass
    else:
        raise AssertionError("block should raise")
    assert "Check your API key" in llm._error(401, '{"error":{"message":"bad key"}}')


class Mock(BaseHTTPRequestHandler):
    def log_message(self, *a):
        pass

    def do_POST(self):
        self.rfile.read(int(self.headers.get("content-length", 0)))
        if "bad" in self.path:
            self.send_response(429)
            self.end_headers()
            self.wfile.write(b'{"error":{"message":"slow down"}}')
            return
        self.send_response(200)
        self.end_headers()
        if self.path.startswith("/v1beta/"):
            lines = [f'data: {json.dumps({"candidates": [{"content": {"parts": [{"text": t}]}}]})}' for t in ("Hel", "lo")]
        elif self.path.endswith("/chat/completions"):
            lines = [f'data: {json.dumps({"choices": [{"delta": {"content": t}}]})}' for t in ("Hel", "lo")] + ["data: [DONE]"]
        else:
            lines = [json.dumps({"message": {"content": t}, "done": False}) for t in ("Hel", "lo")]
        self.wfile.write(("\n".join(lines) + "\n").encode())


@test
def gemini_auto_model_falls_back_when_retired():
    class G(BaseHTTPRequestHandler):
        def log_message(self, *a):
            pass

        def do_GET(self):
            ms = [("gemini-3.6-flash", ["generateContent"]), ("gemini-3-flash-preview", ["generateContent"]),
                  ("gemini-3.6-flash-tts", ["generateContent"]), ("gemini-embedding-001", ["embedContent"])]
            self.send_response(200)
            self.end_headers()
            self.wfile.write(json.dumps({"models": [{"name": f"models/{n}", "supportedGenerationMethods": m} for n, m in ms]}).encode())

        def do_POST(self):
            self.rfile.read(int(self.headers.get("content-length", 0)))
            if "gemini-3.6-flash:" in self.path:  # newest model "retired"
                self.send_response(404)
                self.end_headers()
                self.wfile.write(b'{"error":{"message":"not found"}}')
                return
            self.send_response(200)
            self.end_headers()
            self.wfile.write(f'data: {json.dumps({"candidates": [{"content": {"parts": [{"text": "ok from " + self.path.split("/")[-1].split(":")[0]}]}}]})}\n'.encode())

    srv = HTTPServer(("127.0.0.1", 0), G)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    base = f"http://127.0.0.1:{srv.server_port}"
    llm._models_cache.clear()
    try:
        assert llm.gemini_candidates("k", base) == ("gemini-3.6-flash", "gemini-3-flash-preview")
        out = "".join(llm.stream_chat(llm.LLM("Gemini (free key)", "auto", "k", base), "s", [{"role": "user", "content": "hi"}]))
        assert out == "ok from gemini-3-flash-preview", out
    finally:
        srv.shutdown()


@test
def streaming_against_mock_server():
    srv = HTTPServer(("127.0.0.1", 0), Mock)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    base = f"http://127.0.0.1:{srv.server_port}"
    try:
        msgs = [{"role": "user", "content": "hi"}]
        for cfg in (llm.LLM("Gemini (free key)", "m", "k", base), llm.LLM("OpenAI-compatible (x)", "m", "k", base + "/v1"),
                    llm.LLM("Ollama (local)", "m", "", base)):
            assert "".join(llm.stream_chat(cfg, "sys", msgs)) == "Hello", cfg.provider
        try:
            llm.complete(llm.LLM("OpenAI-compatible (x)", "m", "k", base + "/bad"), "s", msgs)
        except llm.LLMError as e:
            assert "429" in str(e) and "slow down" in str(e)
        else:
            raise AssertionError("expected LLMError")
        try:
            llm.complete(llm.LLM("Gemini (free key)", "m", ""), "s", msgs)
        except llm.LLMError as e:
            assert "key" in str(e).lower()
    finally:
        srv.shutdown()


class RecGemini(BaseHTTPRequestHandler):
    """Gemini mock that records (model, thinkingConfig) for every request."""
    log: list = []

    def log_message(self, *a):
        pass

    def do_GET(self):
        ms = [("gemini-3.6-flash", ["generateContent"]), ("gemini-3.1-flash-lite", ["generateContent"]),
              ("gemini-2.5-flash", ["generateContent"]), ("gemini-3.6-flash-tts", ["generateContent"])]
        self.send_response(200)
        self.end_headers()
        self.wfile.write(json.dumps({"models": [{"name": f"models/{n}", "supportedGenerationMethods": m} for n, m in ms]}).encode())

    def do_POST(self):
        body = json.loads(self.rfile.read(int(self.headers.get("content-length", 0))) or b"{}")
        name = self.path.split("/")[-1].split(":")[0]
        cfg = body.get("generationConfig", {})
        RecGemini.log.append((name, cfg.get("thinkingConfig")))
        if name == "gemini-3.6-flash":
            self.send_response(404)
            self.end_headers()
            self.wfile.write(b'{"error":{"message":"retired"}}')
            return
        if name == "gemini-2.5-flash" and "thinkingConfig" in cfg:
            self.send_response(400)
            self.end_headers()
            self.wfile.write(b'{"error":{"message":"Thinking config unsupported"}}')
            return
        self.send_response(200)
        self.end_headers()
        for piece in ("Hel", "lo"):
            self.wfile.write(f'data: {json.dumps({"candidates": [{"content": {"parts": [{"text": piece}]}}]})}\n\n'.encode())


def _gem_server():
    srv = HTTPServer(("127.0.0.1", 0), RecGemini)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    llm._models_cache.clear()
    llm._good.clear()
    llm._no_think.clear()
    RecGemini.log = []
    return srv, f"http://127.0.0.1:{srv.server_port}"


@test
def gemini_speed_order_thinking_off_and_memory():
    srv, base = _gem_server()
    try:
        assert llm.gemini_candidates("k", base, "Fast")[:2] == ("gemini-3.1-flash-lite", "gemini-3.6-flash")
        assert llm.gemini_candidates("k", base, "Accurate")[0] == "gemini-3.6-flash"
        cfg = llm.LLM("Gemini (free key)", "auto", "k", base, "Accurate")
        info = {}
        out = "".join(llm.stream_chat(cfg, "s", [{"role": "user", "content": "hi"}], info))
        assert out == "Hello" and info["model"] == "gemini-2.5-flash"                   # 3.6 retired, 2.5 next
        assert RecGemini.log[0] == ("gemini-3.6-flash", {"thinkingLevel": "minimal"})   # thinking is off
        RecGemini.log.clear()
        "".join(llm.stream_chat(cfg, "s", [{"role": "user", "content": "again"}]))
        assert RecGemini.log == [("gemini-2.5-flash", None)]   # winner remembered, retired model and rejected setting skipped
    finally:
        srv.shutdown()


@test
def gemini_retries_without_thinking_when_rejected():
    srv, base = _gem_server()
    try:
        cfg = llm.LLM("Gemini (free key)", "gemini-2.5-flash", "k", base)
        assert "".join(llm.stream_chat(cfg, "s", [{"role": "user", "content": "hi"}])) == "Hello"
        assert RecGemini.log == [("gemini-2.5-flash", {"thinkingBudget": 0}), ("gemini-2.5-flash", None)]
    finally:
        srv.shutdown()


if __name__ == "__main__":
    failed = 0
    for fn in TESTS:
        try:
            fn()
            print(f"PASS  {fn.__name__}")
        except Skip as s:
            print(f"SKIP  {fn.__name__} ({s})")
        except Exception as e:
            failed += 1
            print(f"FAIL  {fn.__name__}: {type(e).__name__}: {e}")
    print(f"\n{len(TESTS) - failed} of {len(TESTS)} tests passed or skipped")
    sys.exit(1 if failed else 0)
